"""deckkit v2 — AICA Lab 템플릿(Lab-template.pptx)을 그대로 사용하는 강의 슬라이드 빌더.

템플릿 사용 원칙
  * 표지      = 템플릿 1번 슬라이드 (Title Slide) — 행사명/제목/발표자 개체 틀만 채움
  * 구역 표지 = 템플릿 '제목 슬라이드' 레이아웃 (AICA·SSU 로고 포함)
  * 본문      = 템플릿 'Title and Content' 레이아웃 — 제목 개체 틀 + 본문 개체 틀(템플릿 글머리표 ▪ / -)
                + 템플릿 2번 슬라이드의 날짜·발표자·쪽번호 푸터
  * 마지막    = 템플릿 3번 슬라이드 (Question ?)
  * 색        = 템플릿 헤더 그라데이션/AICA 로고에서 추출한 색만 사용

사용:
    from deckkit import *
    d = Deck(1, "Tensor와 학습", date="2026-10-05")
    s = d.slide("벡터와 행렬", lead="모든 입력은 숫자 배열이 된다", stage="아이디어", notes="...")
    L, R = s.cols(0.55)
    s.bullets(["**벡터**는 숫자의 묶음", ("하위 항목", 1)], L)
    s.image("assets/week1/x.png", R)
    d.save("../lectures/week01.pptx")

인라인 마크업: **굵게**  ==강조(주황)==  `고정폭`   · 좌표 단위 inch
"""
from __future__ import annotations

import copy
import math
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "Lab-template.pptx")

# 템플릿의 행사명 자리(표지 상단 'AICA Lab Meeting', 푸터)에 들어갈 이름
EVENT = "AI Zero 2 Transformer"
AUTHOR = "Yun-Hong Min"

# ---- 템플릿에서 추출한 색 ----
TEAL = "01688F"      # 헤더 그라데이션 시작색 / AICA 로고 진청록
SKY = "72ABC8"       # 헤더 그라데이션 중간색
CYAN = "00A3CA"      # AICA 로고 청색
AQUA = "58C4C4"      # AICA 로고 밝은 청록
ORANGE = "E97132"    # 템플릿 테마 accent2 — 경고·핵심 수치에만
INK = "1B2A36"       # 본문 진한 글자
GRAY = "5E6B75"
LINE = "C5D9E4"
TINT = "EDF5F9"      # SKY의 아주 옅은 틴트 (카드 배경)
TINT2 = "D8EAF2"
ORANGE_TINT = "FDF1E9"
CODE_BG = "F4F7F9"
WHITE = "FFFFFF"

FONT = "맑은 고딕"
MONO = "Courier New"

# 본문 영역
X0, X1 = 0.5, 12.83
W = X1 - X0
Y_TOP = 0.92          # 헤더(0.745") 바로 아래
Y1 = 6.98             # 푸터 위

STAGES = {
    "문제": (ORANGE, WHITE),
    "아이디어": (CYAN, WHITE),
    "계산": (TEAL, WHITE),
    "검증": ("2E8F8F", WHITE),
    "역사": (SKY, WHITE),
    "코드": (INK, WHITE),
    "정리": (TEAL, WHITE),
    "도입": (SKY, WHITE),
    "실습": (INK, WHITE),
}


def rgb(h):
    return RGBColor.from_string(h)


# ================================================================ text utils
_TOKEN = re.compile(r"(\*\*.+?\*\*|==.+?==|`.+?`)")


def parse_inline(text):
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], {"bold": True}))
        elif part.startswith("==") and part.endswith("==") and len(part) > 4:
            out.append((part[2:-2], {"bold": True, "color": ORANGE}))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], {"mono": True}))
        else:
            out.append((part, {}))
    return out


def plain(text):
    return re.sub(r"\*\*|==|`", "", text)


def text_units(s):
    """대략적인 글자 폭 (1.0 = 한글 한 글자 = 글꼴 크기)."""
    u = 0.0
    for ch in s:
        o = ord(ch)
        if 0xAC00 <= o <= 0xD7A3 or 0x3130 <= o <= 0x318F or 0x4E00 <= o <= 0x9FFF:
            u += 1.0
        elif ch == " ":
            u += 0.3
        elif ch in "ilj.,:;|!'`()[]":
            u += 0.32
        elif ch.isupper() or ch in "mwMW@%":
            u += 0.68
        elif o > 0x2000:
            u += 0.8
        else:
            u += 0.56
    return u


def _norm(items):
    if isinstance(items, str):
        items = [items]
    out = []
    for it in items:
        text, lvl = it if isinstance(it, tuple) else (it, 0)
        for line in str(text).split("\n"):
            out.append((line, lvl))
    return out


def est_height(items, w_in, size, spacing=1.2, gap=0.45, sub_ratio=0.85, indent=0.3):
    h = 0.0
    for text, lvl in _norm(items):
        sz = size if lvl == 0 else size * sub_ratio
        avail = max(0.4, w_in - indent * (lvl + 1))
        per_line = avail * 72 / sz
        lines = max(1, math.ceil(text_units(plain(text)) * 1.07 / per_line))
        h += lines * sz * spacing / 72 + gap * sz / 72
    return h + 0.1


def fit(items, w, h, size, min_size, **kw):
    s = size
    while s > min_size and est_height(items, w, s, **kw) > h:
        s -= 0.5
    if est_height(items, w, s, **kw) > h * 1.04:
        WARN.append(f"텍스트 넘침 가능: {plain(str(_norm(items)[0][0]))[:30]}… ({s}pt, box {w:.1f}x{h:.1f})")
    return s


WARN: list[str] = []


def _font(run, size=None, bold=None, color=None, mono=False, italic=False, set_face=True):
    f = run.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic:
        f.italic = True
    if color:
        f.color.rgb = rgb(color)
    if set_face or mono:
        name = MONO if mono else FONT
        f.name = name
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", FONT if tag == "a:ea" else name)


def _clear_bullets(pPr):
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont", "a:buSzPct"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)


def _bullet(p, lvl, color=TEAL):
    """템플릿과 같은 글머리표: 1수준 ▪(Wingdings §), 2수준 -."""
    pPr = p._p.get_or_add_pPr()
    _clear_bullets(pPr)
    ind = Inches(0.26)
    pPr.set("marL", str(int(Inches(0.02 + 0.34 * lvl)) + ind))
    pPr.set("indent", str(-ind))
    c = pPr.makeelement(qn("a:buClr"), {})
    c.append(c.makeelement(qn("a:srgbClr"), {"val": color}))
    pPr.append(c)
    if lvl == 0:
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Wingdings", "pitchFamily": "2", "charset": "2"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "§"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "-"}))


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    _clear_bullets(pPr)
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def fill(tf, items, size=16, color=INK, bold=False, align="l", bullets=False, anchor="t",
         spacing=1.1, gap=0.45, margin=0.08, sub_ratio=0.85, set_face=True, bullet_color=TEAL):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    # 기존 문단 제거
    txBody = tf._txBody
    ps = txBody.findall(qn("a:p"))
    for p in ps[1:]:
        txBody.remove(p)
    p0 = tf.paragraphs[0]
    for r in list(p0._p):
        if r.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0._p.remove(r)
    first = True
    for text, lvl in _norm(items):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        sz = size if lvl == 0 else round(size * sub_ratio * 2) / 2
        p.space_before = Pt(0)
        p.space_after = Pt(sz * gap)
        if bullets:
            _bullet(p, lvl, bullet_color)
        else:
            _no_bullet(p)
        for seg, st in parse_inline(text):
            r = p.add_run()
            r.text = seg
            _font(r, size=sz, bold=st.get("bold", bold), color=st.get("color", color),
                  mono=st.get("mono", False), set_face=set_face)
    return tf


# ================================================================ geometry
class Box:
    def __init__(self, x, y, w, h):
        self.x, self.y, self.w, self.h = x, y, w, h

    @property
    def r(self):
        return self.x + self.w

    @property
    def b(self):
        return self.y + self.h

    def inset(self, dx=0.15, dy=None):
        dy = dx if dy is None else dy
        return Box(self.x + dx, self.y + dy, self.w - 2 * dx, self.h - 2 * dy)

    def cols(self, *ratios, gap=0.3):
        if len(ratios) == 1:
            ratios = (ratios[0], 1 - ratios[0])
        tot = sum(ratios)
        avail = self.w - gap * (len(ratios) - 1)
        out, x = [], self.x
        for r in ratios:
            w = avail * r / tot
            out.append(Box(x, self.y, w, self.h))
            x += w + gap
        return out

    def rows(self, *ratios, gap=0.25):
        if len(ratios) == 1:
            ratios = (ratios[0], 1 - ratios[0])
        tot = sum(ratios)
        avail = self.h - gap * (len(ratios) - 1)
        out, y = [], self.y
        for r in ratios:
            h = avail * r / tot
            out.append(Box(self.x, y, self.w, h))
            y += h + gap
        return out

    def top(self, h, gap=0.25):
        return Box(self.x, self.y, self.w, h), Box(self.x, self.y + h + gap, self.w, self.h - h - gap)

    def bottom(self, h, gap=0.25):
        return Box(self.x, self.y, self.w, self.h - h - gap), Box(self.x, self.b - h, self.w, h)

    def grid(self, n, cols, gap=0.25, row_gap=None):
        row_gap = gap if row_gap is None else row_gap
        rows = math.ceil(n / cols)
        cw = (self.w - gap * (cols - 1)) / cols
        ch = (self.h - row_gap * (rows - 1)) / rows
        return [Box(self.x + (i % cols) * (cw + gap), self.y + (i // cols) * (ch + row_gap), cw, ch)
                for i in range(n)]


# ================================================================ slide context
class S:
    """본문 슬라이드 하나. 템플릿 본문 개체 틀(body)을 bullets()에서 사용한다."""

    def __init__(self, deck, slide, has_lead):
        self.d = deck
        self.s = slide
        self.body_ph = None
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                self.body_ph = ph
        self.area = Box(X0, (Y_TOP + 0.62) if has_lead else (Y_TOP + 0.18), W,
                        Y1 - ((Y_TOP + 0.62) if has_lead else (Y_TOP + 0.18)))
        self._used_body = False

    # -- layout helpers
    def cols(self, *ratios, gap=0.35):
        return self.area.cols(*ratios, gap=gap)

    def rows(self, *ratios, gap=0.25):
        return self.area.rows(*ratios, gap=gap)

    # -- template body placeholder bullets
    def bullets(self, items, box=None, size=18, min_size=12, gap=0.5, anchor="t", color=INK):
        box = box or self.area
        if self._used_body or self.body_ph is None:
            return self.textbox(items, box, size=size, min_size=min_size, bullets=True, gap=gap,
                                anchor=anchor, color=color)
        ph = self.body_ph
        ph.left, ph.top, ph.width, ph.height = Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h)
        sz = fit(items, box.w - 0.2, box.h - 0.1, size, min_size, gap=gap)
        fill(ph.text_frame, items, size=sz, bullets=True, gap=gap, anchor=anchor, color=color,
             set_face=False)
        bp = ph.text_frame._txBody.find(qn("a:bodyPr"))
        for el in list(bp):
            bp.remove(el)
        bp.append(bp.makeelement(qn("a:normAutofit"), {}))
        self._used_body = True
        return ph

    # -- free text
    def textbox(self, items, box, size=16, min_size=10, color=INK, bold=False, align="l",
                bullets=False, anchor="t", gap=0.4, spacing=1.1, margin=0.08, autofit=True):
        if autofit:
            size = fit(items, box.w - 2 * margin - (0.3 if bullets else 0), box.h - 0.08, size,
                       min_size, gap=gap, spacing=spacing + 0.1)
        tb = self.s.shapes.add_textbox(Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
        fill(tb.text_frame, items, size=size, color=color, bold=bold, align=align, bullets=bullets,
             anchor=anchor, gap=gap, spacing=spacing, margin=margin)
        return tb

    def rect(self, box, fill_color=TINT, line=None, radius=0.1, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line_w=1.0, dash=False):
        shp = self.s.shapes.add_shape(shape, Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
        if fill_color:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill_color)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = 4
        else:
            shp.line.fill.background()
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                shp.adjustments[0] = min(0.5, radius / max(0.01, min(box.w, box.h)))
            except Exception:
                pass
        spPr = shp._element.spPr
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
        return shp

    def label(self, shp, items, size=14, color=INK, bold=False, align="c", anchor="m", min_size=9,
              margin=0.1, gap=0.2):
        w = shp.width / 914400 - 2 * margin
        h = shp.height / 914400 - 0.08
        size = fit(items, w, h, size, min_size, gap=gap, spacing=1.2)
        fill(shp.text_frame, items, size=size, color=color, bold=bold, align=align, anchor=anchor,
             margin=margin, gap=gap)
        return shp

    def arrow(self, x1, y1, x2, y2, color=SKY, width=2.0):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        return c

    def badge(self, x, y, d, text, fill_color=TEAL, color=WHITE, size=13):
        c = self.rect(Box(x, y, d, d), fill_color=fill_color, shape=MSO_SHAPE.OVAL)
        self.label(c, str(text), size=size, color=color, bold=True, margin=0.0)
        return c

    # -- composites
    def image(self, path, box, align="c"):
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
        r = min(box.w / iw, box.h / ih)
        pw, ph = iw * r, ih * r
        x = box.x + {"c": (box.w - pw) / 2, "l": 0, "r": box.w - pw}[align]
        return self.s.shapes.add_picture(path, Inches(x), Inches(box.y + (box.h - ph) / 2),
                                         Inches(pw), Inches(ph))

    def card(self, box, head, body=None, tone="teal", body_size=15, head_size=17, bullets=None,
             num=None, min_size=10, fit_h=False, head_fixed=False, head_h=None):
        """카드 하나. fit_h=True면 내용 높이로 줄이고 아래 남은 영역(Box)을 반환."""
        rest = None
        if fit_h:
            need = self.card_need(box.w, head, body, body_size, head_size, num is not None)
            if need < box.h:
                rest = Box(box.x, box.y + need + 0.25, box.w, box.h - need - 0.25)
                box = Box(box.x, box.y, box.w, need)
        fillc, headc, line = {"teal": (TINT, TEAL, None), "accent": (ORANGE_TINT, ORANGE, None),
                              "plain": (WHITE, TEAL, LINE), "dark": (TEAL, WHITE, None)}[tone]
        self.rect(box, fill_color=fillc, line=line)
        bodyc = WHITE if tone == "dark" else INK
        inner = box.inset(0.18, 0.14)
        hx = inner.x
        if num is not None:
            self.badge(inner.x, inner.y + 0.02, 0.42, num, fill_color=headc if tone != "dark" else WHITE,
                       color=WHITE if tone != "dark" else TEAL, size=12)
            hx = inner.x + 0.55
        hh = head_h or min(0.9, max(0.42, est_height([head], inner.r - hx, head_size, gap=0) - 0.05))
        self.textbox(head, Box(hx, inner.y, inner.r - hx, hh), size=head_size, bold=True, color=headc,
                     anchor="m", min_size=11, margin=0.02, gap=0, autofit=not head_fixed)
        if body:
            by = inner.y + hh + 0.08
            items = body
            if bullets is None:
                bullets = isinstance(body, list) and len(body) > 1
            self.textbox(items, Box(inner.x, by, inner.w, inner.b - by), size=body_size,
                         bullets=bullets, color=bodyc, min_size=min_size, margin=0.02, gap=0.35)
        return rest

    def card_need(self, w, head, body, body_size=15, head_size=17, num=False):
        iw = w - 0.36
        hw = iw - (0.55 if num else 0)
        hh = min(0.9, max(0.42, est_height([head], hw, head_size, gap=0) - 0.05))
        bh = 0
        if body:
            bl = isinstance(body, list) and len(body) > 1
            bh = est_height(body, iw - (0.3 if bl else 0) - 0.04, body_size, gap=0.35, spacing=1.2) + 0.08
        return 0.28 + hh + bh + 0.1

    def cards(self, cards, box=None, cols=None, numbered=False, body_size=15, head_size=17, gap=0.25,
              fit_h=True, min_h=1.2):
        """카드 격자. fit_h=True면 내용 높이에 맞춰 줄이고, 남은 영역(Box)을 반환."""
        box = box or self.area
        n = len(cards)
        cols = cols or (n if n <= 4 else 3)
        rows = math.ceil(n / cols)
        cw = (box.w - gap * (cols - 1)) / cols
        full_h = (box.h - gap * (rows - 1)) / rows
        heights = []
        for r in range(rows):
            need = max(self.card_need(cw, c["head"], c.get("body"), body_size, head_size, numbered or c.get("num"))
                       for c in cards[r * cols:(r + 1) * cols])
            heights.append(min(full_h, max(min_h, need)) if fit_h else full_h)
        # 같은 격자 안에서는 제목 글자 크기를 통일
        hs = head_size
        for c in cards:
            hw = cw - 0.36 - (0.55 if (numbered or c.get("num")) else 0)
            while hs > 11 and est_height([c["head"]], hw, hs, gap=0) - 0.05 > 0.9:
                hs -= 0.5
        head_size = hs
        y = box.y
        for r in range(rows):
            row = cards[r * cols:(r + 1) * cols]
            hh = max(min(0.9, max(0.42, est_height([c["head"]], cw - 0.36 - (0.55 if (numbered or c.get("num")) else 0),
                                                   head_size, gap=0) - 0.05)) for c in row)
            for k, c in enumerate(row):
                i = r * cols + k
                b = Box(box.x + k * (cw + gap), y, cw, heights[r])
                self.card(b, c["head"], c.get("body"), tone=c.get("tone", "teal"), body_size=body_size,
                          head_size=head_size, num=(i + 1) if numbered else c.get("num"),
                          bullets=c.get("bullets"), head_fixed=True, head_h=hh)
            y += heights[r] + gap
        return Box(box.x, y + 0.05, box.w, max(0.0, box.b - y - 0.05))

    def flow(self, steps, box, body_size=13, head_size=16, gap=0.42, vertical=False):
        """steps: [{'head','body','tone'}] 사이를 화살표로 연결."""
        n = len(steps)
        if vertical:
            bh = (box.h - gap * (n - 1)) / n
            for i, st in enumerate(steps):
                b = Box(box.x, box.y + i * (bh + gap), box.w, bh)
                self._flow_box(b, st, body_size, head_size)
                if i < n - 1:
                    self.arrow(b.x + b.w / 2, b.b + 0.04, b.x + b.w / 2, b.b + gap - 0.04)
            return
        bw = (box.w - gap * (n - 1)) / n
        for i, st in enumerate(steps):
            b = Box(box.x + i * (bw + gap), box.y, bw, box.h)
            self._flow_box(b, st, body_size, head_size)
            if i < n - 1:
                self.arrow(b.r + 0.05, b.y + b.h / 2, b.r + gap - 0.05, b.y + b.h / 2)

    def _flow_box(self, b, st, body_size, head_size):
        st = st if isinstance(st, dict) else {"head": st}
        tone = st.get("tone", "teal")
        fillc, headc, line = {"teal": (TINT, TEAL, None), "accent": (ORANGE_TINT, ORANGE, ORANGE),
                              "plain": (WHITE, TEAL, LINE), "dark": (TEAL, WHITE, None)}[tone]
        self.rect(b, fill_color=fillc, line=line)
        inner = b.inset(0.1, 0.1)
        if st.get("body"):
            hh = min(0.8, max(0.4, est_height([st["head"]], inner.w, head_size, gap=0)))
            self.textbox(st["head"], Box(inner.x, inner.y, inner.w, hh), size=head_size, bold=True,
                         color=headc, align="c", anchor="m", min_size=10, margin=0.02, gap=0)
            self.textbox(st["body"], Box(inner.x, inner.y + hh + 0.05, inner.w, inner.h - hh - 0.05),
                         size=body_size, align="c", min_size=9, margin=0.02, gap=0.25,
                         color=WHITE if tone == "dark" else INK, anchor="m" if b.h > 2.4 else "t")
        else:
            self.textbox(st["head"], inner, size=head_size, bold=True, color=headc, align="c",
                         anchor="m", min_size=10, margin=0.02, gap=0)

    def formula(self, lines, box, size=24, min_size=13, fill_color=CODE_BG):
        lines = [lines] if isinstance(lines, str) else lines
        self.rect(box, fill_color=fill_color, line=LINE)
        def mono_units(t):
            return sum(1.0 if (0xAC00 <= ord(c) <= 0xD7A3) else 0.61 for c in plain(t))
        longest = max(mono_units(l) for l in lines)
        s = size
        while s > min_size and (longest * s / 72 > box.w - 0.45 or len(lines) * s * 1.35 / 72 > box.h - 0.1):
            s -= 0.5
        tb = self.s.shapes.add_textbox(Inches(box.x + 0.15), Inches(box.y), Inches(box.w - 0.3), Inches(box.h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            _no_bullet(p)
            for seg, st in parse_inline(ln):
                r = p.add_run()
                r.text = seg
                _font(r, size=s, bold=True, color=st.get("color", INK), mono=True)
        return tb

    def symbols(self, pairs, box, size=15, key_w=1.7):
        """[(기호, 뜻)] 표 형태 목록."""
        n = len(pairs)
        rh = min(0.62, box.h / n)
        for i, (k, v) in enumerate(pairs):
            y = box.y + i * rh
            chip = self.rect(Box(box.x, y + 0.05, key_w, rh - 0.1), fill_color=TINT2)
            self.label(chip, f"`{k}`" if "`" not in k else k, size=size, bold=True, margin=0.04)
            self.textbox(v, Box(box.x + key_w + 0.12, y, box.w - key_w - 0.12, rh), size=size,
                         anchor="m", min_size=10, gap=0)

    def callout(self, text, box, kind="key", size=16, fit_h=True):
        """kind: key(청록 채움) | tip(옅은 청록) | warn(주황). fit_h면 내용 높이로 줄인다."""
        if fit_h:
            need = est_height([text], box.w - 0.5, size, gap=0.2, spacing=1.25) + 0.25
            if need < box.h:
                box = Box(box.x, box.y, box.w, max(0.6, need))
        self.last_callout_box = box
        fillc, color, line = {"key": (TEAL, WHITE, None), "tip": (TINT, INK, SKY),
                              "warn": (ORANGE_TINT, INK, ORANGE)}[kind]
        b = self.rect(box, fill_color=fillc, line=line)
        self.label(b, text, size=size, color=color, bold=(kind == "key"), align="c" if kind == "key" else "l",
                   margin=0.2)
        return b

    def takeaway(self, text, size=16):
        """영역 맨 아래 한 줄 결론 바. 영역을 줄여 준다."""
        b = Box(X0, Y1 - 0.58, W, 0.58)
        self.callout(text, b, kind="key", size=size)
        self.area = Box(self.area.x, self.area.y, self.area.w, self.area.h - 0.8)
        return b

    def table(self, header, rows, box=None, widths=None, size=14, min_size=9, highlight=(),
              first_col_bold=True, align=None):
        box = box or self.area
        nr, nc = len(rows) + 1, len(header)
        widths = widths or [1] * nc
        tot = sum(widths)
        widths = [w * box.w / tot for w in widths]

        def row_h(cells, sz):
            return max(est_height([str(c)], cw - 0.16, sz, spacing=1.2, gap=0) for c, cw in zip(cells, widths)) + 0.06

        sz = size
        while sz > min_size and sum(row_h(r, sz) for r in [header] + rows) > box.h:
            sz -= 0.5
        heights = [row_h(r, sz) for r in [header] + rows]
        if sum(heights) > box.h * 1.05:
            WARN.append(f"표 넘침 가능: {header[0]} ({sz}pt)")
        gt = self.s.shapes.add_table(nr, nc, Inches(box.x), Inches(box.y), Inches(box.w), Inches(sum(heights)))
        tbl = gt.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("firstRow", "1")
        tblPr.set("bandRow", "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
        for j, cw in enumerate(widths):
            tbl.columns[j].width = Inches(cw)
        hl = set(highlight)
        for i, r in enumerate([header] + rows):
            tbl.rows[i].height = Inches(heights[i])
            for j, val in enumerate(r):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                if i == 0:
                    cell.fill.fore_color.rgb = rgb(TEAL)
                elif (i - 1) in hl:
                    cell.fill.fore_color.rgb = rgb(ORANGE_TINT)
                else:
                    cell.fill.fore_color.rgb = rgb(WHITE if i % 2 else TINT)
                a = "c" if i == 0 else (align[j] if align else "l")
                fill(cell.text_frame, [str(val)], size=sz, color=WHITE if i == 0 else INK,
                     bold=(i == 0 or (j == 0 and first_col_bold)), align=a, margin=0.04, gap=0)
                _cell_border(cell)
        return gt

    def code(self, code, box, size=14, min_size=9, fit_h=True):
        lines = code.rstrip("\n").split("\n")
        longest = max(len(l) for l in lines)
        sz = size
        while sz > min_size and (len(lines) * sz * 1.22 / 72 > box.h - 0.3 or longest * sz * 0.61 / 72 > box.w - 0.3):
            sz -= 0.5
        if fit_h:
            box = Box(box.x, box.y, box.w, min(box.h, len(lines) * sz * 1.22 / 72 + 0.4))
        self.last_code_box = box
        self.rect(box, fill_color=CODE_BG, line=LINE, radius=0.06)
        tb = self.s.shapes.add_textbox(Inches(box.x + 0.12), Inches(box.y + 0.12), Inches(box.w - 0.24),
                                       Inches(box.h - 0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _no_bullet(p)
            p.line_spacing = 1.0
            code_part, _, comment = ln.partition("#")
            if code_part:
                r = p.add_run()
                r.text = code_part
                _font(r, size=sz, color=INK, mono=True)
            if _:
                r = p.add_run()
                r.text = "#" + comment
                _font(r, size=sz, color="3F8A6E", mono=True)
            if not ln:
                r = p.add_run()
                r.text = " "
                _font(r, size=sz, mono=True)
        return tb

    def timeline(self, events, box, highlight=(), body_size=13, fit_h=True):
        """events: [(연도, 제목, 설명)]"""
        n = len(events)
        gap = 0.16
        cw = (box.w - gap * (n - 1)) / n
        need_max = max(0.3 + 0.82 + est_height([e[2]], cw - 0.2, body_size, gap=0.25, spacing=1.2) for e in events)
        line_y = box.y + 0.72
        self.rect(Box(box.x, line_y - 0.025, box.w, 0.05), fill_color=LINE, shape=MSO_SHAPE.RECTANGLE)
        hl = set(highlight)
        for i, (yr, head, body) in enumerate(events):
            x = box.x + i * (cw + gap)
            col = ORANGE if i in hl else TEAL
            self.textbox(str(yr), Box(x, box.y, cw, 0.5), size=20, bold=True, color=col, align="c",
                         anchor="b", min_size=11, margin=0.0, gap=0)
            self.rect(Box(x + cw / 2 - 0.12, line_y - 0.12, 0.24, 0.24), fill_color=col, shape=MSO_SHAPE.OVAL)
            need = 0.2 + 0.82 + est_height([body] if isinstance(body, str) else body, cw - 0.2, body_size, gap=0.25, spacing=1.2)
            cb = Box(x, line_y + 0.3, cw, box.b - line_y - 0.3)
            if fit_h:
                cb = Box(x, line_y + 0.3, cw, min(cb.h, max(need_max, 1.4)))
            self.rect(cb, fill_color=ORANGE_TINT if i in hl else TINT)
            inner = cb.inset(0.08, 0.1)
            self.textbox(head, Box(inner.x, inner.y, inner.w, 0.75), size=15, bold=True, color=col,
                         align="c", anchor="m", min_size=10, margin=0.02, gap=0)
            self.textbox(body, Box(inner.x, inner.y + 0.82, inner.w, inner.h - 0.82), size=body_size,
                         min_size=9, margin=0.02, gap=0.25)


def _cell_border(cell, color=LINE):
    tcPr = cell._tc.get_or_add_tcPr()
    fills = tcPr.findall(qn("a:solidFill"))
    for f in fills:
        tcPr.remove(f)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {"w": "6350"})
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(sf)
        tcPr.append(ln)
    for f in fills:
        tcPr.append(f)


# ================================================================ deck
COURSE = [
    (1, "배울 수\n있을까?", "Tensor · 학습", "1943–1958"),
    (2, "직선으로\n안 되면?", "Perceptron · MLP", "1958–1986"),
    (3, "공간과\n깊이는?", "CNN · ResNet", "1989–2016"),
    (4, "먼 정보를\n기억하려면?", "RNN · LSTM", "1986–2014"),
    (5, "원문을\n다시 보면?", "Attention", "2014–2017"),
    (6, "참조만으로\n만들면?", "Transformer", "2017–"),
]


class Deck:
    def __init__(self, week, title, date, event=EVENT):
        self.prs = Presentation(TEMPLATE)
        self.week = week
        self.date = date
        self.event = event
        s_title, s_proto, s_last = list(self.prs.slides)
        self.layout = s_proto.slide_layout
        self.sec_layout = next(l for l in self.prs.slide_masters[1].slide_layouts if l.name == "제목 슬라이드")
        # --- 표지: 템플릿 개체 틀 그대로, 글자만 교체
        for sh in s_title.placeholders:
            idx = sh.placeholder_format.idx
            if idx == 10:
                _replace_text(sh, event)
            elif idx == 0:
                _replace_text(sh, f"Week {week}. {title}")
                n = text_units(f"Week {week}. {title}")
                if n > 24:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(max(22, int(32 * 24 / n)))
            # idx 1 (발표자/이메일)은 템플릿 그대로
        # --- 푸터 원본(날짜·발표자·쪽번호)
        self._footer = [copy.deepcopy(sp._element) for sp in s_proto.shapes
                        if not (sp.is_placeholder and sp.placeholder_format.idx in (0, 1))]
        self._ctxs = []
        self.n = 0

    # ---------------------------------------------------------------- slides
    def slide(self, title, lead=None, stage=None, notes=None):
        s = self.prs.slides.add_slide(self.layout)
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 0:
                tf = ph.text_frame
                tf.text = title
                n = text_units(title)
                if n > 30:
                    for r in tf.paragraphs[0].runs:
                        r.font.size = Pt(max(20, int(32 * 30 / n)))
        tree = s.shapes._spTree
        for el in self._footer:
            el = copy.deepcopy(el)
            for t in el.iter(qn("a:t")):
                if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t.text or ""):
                    t.text = self.date
                elif (t.text or "").strip() == "AICA Lab Meeting":
                    t.text = self.event
            if el.find(".//" + qn("p:ph")) is None:  # 발표자 텍스트 상자: 두 자리 쪽번호와 겹치지 않게 왼쪽으로
                off = el.find(".//" + qn("a:off"))
                if off is not None:
                    off.set("x", str(int(off.get("x")) - int(Inches(0.3))))
            tree.append(el)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        ctx = S(self, s, has_lead=bool(lead) or bool(stage))
        if lead:
            lw = W - (1.55 if stage else 0)
            ctx.textbox(lead, Box(X0, Y_TOP, lw, 0.55), size=19, bold=True, color=TEAL, anchor="m",
                        min_size=13, margin=0.02, gap=0)
        if stage:
            fc, tc = STAGES[stage]
            chip = ctx.rect(Box(X1 - 1.3, Y_TOP + 0.08, 1.3, 0.4), fill_color=fc, radius=0.2)
            fill(chip.text_frame, [stage], size=13, color=tc, bold=True, align="c", anchor="m", margin=0.02, gap=0)
        self.n += 1
        ctx._base_n = len(s.shapes)
        self._ctxs.append(ctx)
        return ctx

    def finish(self, ctx):
        """본문 개체 틀을 쓰지 않은 슬라이드에서 빈 개체 틀 제거하고, 본문이 위쪽에만 몰려 있으면 아래로 조금 내려 균형을 맞춘다."""
        if not ctx._used_body and ctx.body_ph is not None:
            ctx.body_ph._element.getparent().remove(ctx.body_ph._element)
            ctx.body_ph = None
            ctx._base_n -= 1  # 개체 틀이 빠졌으므로 본문 도형 시작 위치도 하나 앞으로
        shapes = list(ctx.s.shapes)[ctx._base_n:]
        if ctx._used_body and ctx.body_ph is not None:
            shapes.append(ctx.body_ph)
        if not shapes:
            return
        bottom = max((sh.top + sh.height) / 914400 for sh in shapes)
        gap = Y1 - bottom
        if gap > 0.6:
            dy = Inches(min(gap * 0.5, 0.9))
            for sh in shapes:
                sh.top = sh.top + dy

    def section(self, title, subtitle="", notes=None):
        """템플릿 '제목 슬라이드' 레이아웃(로고 포함)으로 만든 구역 표지."""
        s = self.prs.slides.add_slide(self.sec_layout)
        for ph in s.placeholders:
            idx = ph.placeholder_format.idx
            tf = ph.text_frame
            if idx == 0:
                lines = title.split("\n")
                fill(tf, lines, size=40 if len(lines) == 1 else 36, bold=True, color=INK, align="c",
                     anchor="b", gap=0.1, set_face=True)
                if len(lines) > 1:
                    for r in tf.paragraphs[0].runs:
                        r.font.size = Pt(22)
                        r.font.color.rgb = rgb(TEAL)
            elif idx == 1:
                if subtitle:
                    fill(tf, subtitle.split("\n"), size=20, color=GRAY, align="c", gap=0.2)
                else:
                    ph._element.getparent().remove(ph._element)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.n += 1
        return s

    # ---------------------------------------------------------------- recipes
    def course_map(self, ctx, box, current):
        n = len(COURSE)
        gap = 0.22
        bw = (box.w - gap * (n - 1)) / n
        for i, (wk, q, topic, era) in enumerate(COURSE):
            b = Box(box.x + i * (bw + gap), box.y, bw, box.h)
            cur = wk == current
            done = wk < current
            ctx.rect(b, fill_color=TEAL if cur else (TINT2 if done else TINT),
                     line=None if cur else (None if done else LINE))
            inner = b.inset(0.08, 0.1)
            col = WHITE if cur else (TEAL if done else GRAY)
            h = inner.h
            ctx.textbox(f"Week {wk}", Box(inner.x, inner.y, inner.w, h * 0.17), size=13, bold=True,
                        color=WHITE if cur else SKY, align="c", margin=0, gap=0, anchor="m", autofit=False)
            ctx.textbox(q, Box(inner.x, inner.y + h * 0.18, inner.w, h * 0.42), size=15, bold=True, color=col,
                        align="c", anchor="m", min_size=10, margin=0, gap=0)
            ctx.textbox(topic, Box(inner.x, inner.y + h * 0.61, inner.w, h * 0.22), size=12, color=col,
                        align="c", anchor="m", min_size=8, margin=0, gap=0)
            ctx.textbox(era, Box(inner.x, inner.y + h * 0.83, inner.w, h * 0.17), size=11,
                        color=WHITE if cur else GRAY, align="c", anchor="m", margin=0, gap=0, autofit=False)
            if i < n - 1:
                ctx.arrow(b.r + 0.02, b.y + b.h / 2, b.r + gap - 0.02, b.y + b.h / 2, width=1.5)

    def quiz(self, title, qa, lead="답을 자기 말로 설명할 수 있는가? (정답은 발표자 노트)", notes=""):
        """셀프 체크: 질문은 슬라이드, 정답은 발표자 노트."""
        ans = "\n".join(f"Q{i+1}. {plain(q)}\n  → {plain(a)}" for i, (q, a) in enumerate(qa))
        ctx = self.slide(title, lead=lead, stage="검증", notes=(notes + "\n\n" if notes else "") + "[정답]\n" + ans)
        n = len(qa)
        gap = 0.16
        rh = min(1.05, (ctx.area.h - 0.2 - gap * (n - 1)) / n)
        for i, (q, _) in enumerate(qa):
            b = Box(ctx.area.x, ctx.area.y + i * (rh + gap), ctx.area.w, rh)
            ctx.rect(b, fill_color=TINT)
            dd = min(0.5, rh - 0.12)
            ctx.badge(b.x + 0.15, b.y + (rh - dd) / 2, dd, f"Q{i+1}", size=11)
            ctx.textbox(q, Box(b.x + 0.8, b.y, b.w - 0.95, rh), size=17, anchor="m", min_size=11, gap=0)
        self.finish(ctx)
        return ctx

    # ---------------------------------------------------------------- 공통 구성
    @staticmethod
    def _join(items):
        items = [items] if isinstance(items, str) else items
        return " / ".join(plain(str(i[0] if isinstance(i, tuple) else i)) for i in items)

    def question(self, q, sub="", notes=None):
        """이번 주의 질문 (템플릿 구역 슬라이드)."""
        return self.section(f"Week {self.week} · 이번 주의 질문\n{q}", sub, notes=notes)

    def part(self, k, title, question, notes=None):
        return self.section(f"Part {k}\n{title}", question, notes=notes)

    def bridge(self, learned, remaining, question, notes=None, title="지난 주에서 이번 주로"):
        """지난 주 배운 것 → 남은 문제 → 이번 주 질문."""
        notes = notes or (f"지난 주 복습 — {self._join(learned)}. 하지만 남은 문제가 있다: {self._join(remaining)}. "
                          f"이번 주는 이 질문에서 출발한다: {self._join(question)}. 학생에게 지난 주 내용을 한 문장씩 먼저 말해 보게 한다.")
        ctx = self.slide(title, lead="매 주는 지난 주가 남긴 문제에서 출발한다", stage="도입", notes=notes)
        ctx.flow([{"head": "지난 주에 배운 것", "body": learned},
                  {"head": "아직 풀리지 않은 문제", "body": remaining, "tone": "accent"},
                  {"head": "이번 주의 질문", "body": question, "tone": "dark"}],
                 Box(ctx.area.x, ctx.area.y + 0.3, ctx.area.w, min(ctx.area.h - 0.6, 3.5)),
                 body_size=19, head_size=20, gap=0.55)
        return ctx

    def roadmap(self, parts, goals, notes=None):
        """과정 지도(현재 주 강조) + 오늘의 순서 + 학습 목표."""
        notes = notes or (f"과정 지도에서 오늘은 {self.week}번째 질문이다. 오늘의 순서: {self._join(parts)}. "
                          f"끝나면 할 수 있어야 하는 것: {self._join(goals)}. 수업 마지막에 이 목표를 셀프 체크로 다시 확인한다.")
        ctx = self.slide("오늘의 지도", lead=f"여섯 질문 중 {self.week}번째 — 오늘의 순서와 목표",
                         stage="도입", notes=notes)
        top, bot = ctx.area.top(2.05, gap=0.3)
        self.course_map(ctx, top, self.week)
        L, R = bot.cols(0.5, gap=0.4)
        ctx.rect(L, fill_color=TINT)
        ctx.textbox("오늘의 순서", Box(L.x + 0.2, L.y + 0.1, L.w - 0.4, 0.4), size=16, bold=True,
                    color=TEAL, anchor="m", gap=0)
        items = [f"**Part {i+1}** {t}" for i, t in enumerate(parts)]
        ctx.textbox(items, Box(L.x + 0.2, L.y + 0.55, L.w - 0.4, L.h - 0.65), size=15, min_size=10, gap=0.3)
        ctx.rect(R, fill_color=WHITE, line=LINE)
        ctx.textbox("끝나면 할 수 있는 것", Box(R.x + 0.2, R.y + 0.1, R.w - 0.4, 0.4), size=16, bold=True,
                    color=TEAL, anchor="m", gap=0)
        ctx.textbox(goals, Box(R.x + 0.2, R.y + 0.55, R.w - 0.4, R.h - 0.65), size=15, min_size=10,
                    bullets=True, gap=0.3)
        return ctx

    def glossary(self, rows, notes=None):
        """오늘 새로 나오는 용어: (용어, 한 줄 뜻, 비유/예)."""
        notes = notes or ("오늘 처음 나오는 용어들. 외울 필요는 없고, 설명 중 막히면 이 표로 돌아온다. "
                          + " / ".join(f"{plain(r[0])}: {plain(r[1])}" for r in rows))
        ctx = self.slide("오늘의 새 용어", lead="낯선 단어를 먼저 한 번 보고 시작한다", stage="도입", notes=notes)
        ctx.table(["용어", "한 줄 뜻", "비유 · 예"], rows, widths=[2.2, 5.2, 4.9], size=15)
        return ctx

    def summary(self, items, notes=None, title="한 장 요약"):
        notes = notes or ("오늘의 요약. 각 줄을 학생이 자기 말로 다시 설명하게 한다: " + self._join(items))
        ctx = self.slide(title, lead="오늘 배운 것을 한 줄씩", stage="정리", notes=notes)
        n = len(items)
        gap = 0.14
        rh = min(0.95, (ctx.area.h - gap * (n - 1)) / n)
        for i, it in enumerate(items):
            y = ctx.area.y + i * (rh + gap)
            dd = min(0.5, rh - 0.12)
            ctx.badge(ctx.area.x, y + (rh - dd) / 2, dd, i + 1, size=13)
            ctx.textbox(it, Box(ctx.area.x + dd + 0.25, y, ctx.area.w - dd - 0.3, rh), size=18,
                        anchor="m", min_size=11, gap=0)
        return ctx

    def misconceptions(self, rows, notes=None):
        notes = (notes + " " if notes else "") + ("각 문장이 왜 틀렸는지 계산으로 설명하게 한다. "
                 + " / ".join(f"{plain(a)} → {plain(b)}" for a, b in rows))
        ctx = self.slide("흔한 오해 바로잡기", lead="틀린 문장을 계산의 언어로 고쳐 쓴다", stage="검증", notes=notes)
        ctx.table(["자주 듣는 말", "더 정확한 설명"], rows, widths=[4.6, 7.7], size=15)
        return ctx

    def homework(self, items, notes=None, lab=None):
        notes = (notes + " " if notes else "") + ("과제는 ‘논문 읽고 정리’ 대신 제출 결과가 분명한 질문으로 준다(원자료 운영안). "
                 + " / ".join(f"{plain(h)}: {plain(b)}" for h, b in items))
        ctx = self.slide("이번 주 과제", lead="제출물이 분명한 과제 — 결과를 보기 전에 먼저 예측해 적는다",
                         stage="실습", notes=notes)
        if lab:
            L, R = ctx.cols(0.55)
            ctx.cards([{"head": h, "body": b} for h, b in items], L, cols=1, numbered=True, body_size=14)
            ctx.code(lab, R, size=13)
        else:
            ctx.cards([{"head": h, "body": b} for h, b in items], cols=1, numbered=True, body_size=15)
        return ctx

    def references(self, rows, notes=None):
        notes = notes or ("번호는 원자료 참고문헌 번호. 처음부터 모든 실험을 읽기보다 표에 적은 절과 그림만 먼저 읽는다. "
                          + " / ".join(plain(r[0]) for r in rows))
        ctx = self.slide("참고 자료", lead="원문은 필요한 절만 골라 읽는다", stage="정리", notes=notes)
        ctx.table(["자료", "이번 주에 읽을 부분"], rows, widths=[7.3, 5.0], size=14)
        return ctx

    def handoff(self, solved, remaining, next_q, notes=None):
        """이번 주가 해결한 것 / 남긴 문제 / 다음 주 질문."""
        notes = (notes + " " if notes else "") + (f"이번 주에 해결한 것: {self._join(solved)}. 남은 문제: {self._join(remaining)}. "
                 f"다음: {self._join(next_q)}.")
        ctx = self.slide("남은 문제와 다음 질문", lead="답은 새 문제를 남긴다 — 다음 주는 여기서 출발한다",
                         stage="정리", notes=notes)
        ctx.flow([{"head": "이번 주에 해결한 것", "body": solved},
                  {"head": "아직 남은 문제", "body": remaining, "tone": "accent"},
                  {"head": f"Week {self.week + 1}의 질문" if self.week < 6 else "다음 여정", "body": next_q,
                   "tone": "dark"}], Box(ctx.area.x, ctx.area.y + 0.3, ctx.area.w, min(ctx.area.h - 0.6, 3.5)),
                 body_size=19, head_size=20, gap=0.55)
        return ctx

    def save(self, path):
        for c in self._ctxs:
            self.finish(c)
        prs = self.prs
        lst = prs.slides._sldIdLst
        ids = list(lst)
        proto, last = ids[1], ids[2]
        prs.part.drop_rel(proto.rId)
        lst.remove(proto)
        lst.remove(last)
        lst.append(last)
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        prs.save(path)
        if WARN:
            print("\n".join("  ! " + w for w in WARN))
        return path


def _replace_text(sh, text):
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.add_run().text = text
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


# ================================================================ matplotlib
def mpl_setup():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    f = "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"
    if os.path.exists(f):
        font_manager.fontManager.addfont(f)
        plt.rcParams["font.family"] = "NanumGothic"
    plt.rcParams.update({"axes.unicode_minus": False, "axes.edgecolor": "#" + LINE,
                         "axes.labelcolor": "#" + INK, "xtick.color": "#" + GRAY,
                         "ytick.color": "#" + GRAY, "savefig.dpi": 200, "savefig.bbox": "tight",
                         "axes.spines.top": False, "axes.spines.right": False,
                         "mathtext.fontset": "dejavusans"})
    return plt


HEX = {k: "#" + v for k, v in dict(TEAL=TEAL, SKY=SKY, CYAN=CYAN, AQUA=AQUA, ORANGE=ORANGE, INK=INK,
                                   GRAY=GRAY, LINE=LINE, TINT=TINT, TINT2=TINT2).items()}

# 그림 스크립트(v1)용 호환 팔레트 — 템플릿 색으로 매핑
PALETTE_HEX = {"TEAL": "#" + TEAL, "TEAL2": "#2A93B5", "MINT": "#E3F0F6", "DARK": "#" + INK,
               "GRAY": "#" + GRAY, "ACCENT": "#" + ORANGE, "LINE": "#" + LINE}
